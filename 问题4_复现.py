"""Unique Q4 pipeline. --output-root isolates independent reproduction."""
from __future__ import annotations
import argparse,json,sys,time,shutil,subprocess,importlib.metadata
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
import 问题4_求解 as q

ROOT=q.ROOT
SKILL=Path(r'C:\Users\23572\.codex\skills\math-modeling')

def difference(s,t):
    times=np.intersect1d(s['time_s'],t['time_s'])
    times=times[(times>0)&(times%21600==0)]
    sc=q.sample(s);tc=q.sample(t)
    si=np.searchsorted(s['time_s'],times);ti=np.searchsorted(t['time_s'],times)
    return float(np.nanmax(np.abs(sc[si]-tc[ti])))

def validate_run(s):
    m=s['meta']
    assert m['previous_Cmax']>=.15>m['final_Cmax']
    assert m['moisture_balance_relative']<=1e-8,m
    assert m['heat_balance_scaled']<=1e-12
    assert max(m['heat_linear_residual'],m['moisture_linear_residual'])<=1e-12
    assert max(m['final_matrix_conditions'].values())<=1e8
    assert m['min_moisture']>=0 and m['min_temperature_c']>=28-1e-8 and m['max_temperature_c']<=50.5
    assert m['min_mapping_Jacobian_m']>0
    assert m['max_radial_moisture_increase']<1e-8 and m['max_step_Cmax_increase']<1e-8

def write_outputs(base,runs,out,verification):
    res=out/'results';res.mkdir(parents=True,exist_ok=True)
    t=base['time_s'];R=base['radius_m'];C=base['moisture'];T=base['temperature_c'];x=np.linspace(0,1,C.shape[1])
    arr=q.sample(base,.001)
    wb=openpyxl.load_workbook(q.old.TEMPLATE_DIR/'result4.xlsx');ws=wb.active
    # Template contains illustrative ellipses, expanded into actual data columns.
    for row in ws:
        for cell in row:cell.value=None
    headers=['时间\\到药材中心的距离']+[round(j*.1,1) for j in range(20)]+['药材表面']
    for j,v in enumerate(headers,1):ws.cell(1,j,v)
    for i,tt in enumerate(t[1:],2):
        ws.cell(i,1,float(tt))
        for j,v in enumerate(arr[i-1],2):
            cell=ws.cell(i,j,None if np.isnan(v) else float(v));cell.number_format='0.0000'
    ws.freeze_panes='B2';wb.save(res/'result4.xlsx');wb.close()
    # Re-open and verify every value, mask and numeric format.
    wb=openpyxl.load_workbook(res/'result4.xlsx',data_only=True);ws=wb.active
    assert ws.max_row==len(t) and ws.max_column==22
    for i in range(2,len(t)+1):
        assert ws.cell(i,1).value==t[i-1]
        for j,expected in enumerate(arr[i-1],2):
            val=ws.cell(i,j).value
            assert (val is None) if np.isnan(expected) else (abs(val-expected)<1e-12 and ws.cell(i,j).number_format=='0.0000')
    wb.close()
    mask=(t>0)&((t%21600==0)|(t==t[-1]));table=q.sample(base)[mask]
    df=pd.DataFrame(table,columns=['r0_cm','r0.5_cm','r1_cm','r1.5_cm','surface'])
    df.insert(0,'radius_cm',R[mask]*100);df.insert(0,'time_h',t[mask]/3600);df.insert(0,'time_s',t[mask])
    df.to_csv(res/'问题4_表6含水率.csv',index=False,encoding='utf-8-sig')
    process=pd.DataFrame({'time_s':t,'radius_cm':R*100,'Cmax':C.max(axis=1),'Ccenter':C[:,0],'Csurface':C[:,-1],'Tcenter_C':T[:,0],'Tsurface_C':T[:,-1]})
    process.to_csv(res/'问题4_关键过程.csv',index=False)
    rows=[]
    for name,s in runs.items():
        m=s['meta'];rows.append({'scenario':name,'N':m['N'],'dt_s':m['dt_s'],'a':m['a'],'fixed':m['fixed'],'dry_time_s':m['dry_time_s'],'dry_time_h':m['dry_time_h'],'relative_to_main_percent':100*(m['dry_time_s']/t[-1]-1),'radius_end_cm':m['radius_end_cm'],'runtime_s':m['runtime_s']})
    pd.DataFrame(rows).to_csv(res/'问题4_情景汇总.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame([{'scenario':k,**v} for k,v in verification.items()]).to_csv(res/'问题4_数值加密.csv',index=False)
    unique,counts=np.unique(base['iterations'],return_counts=True)
    pd.DataFrame({'iterations':unique,'step_count':counts}).to_csv(res/'问题4_迭代统计.csv',index=False)
    fixed=runs['fixed']['meta']['dry_time_s']
    etaR=100*(t[-1]-fixed)/fixed
    etaA={str(a):100*(runs[f'shape_{a}']['meta']['dry_time_s']/t[-1]-1) for a in [.1,.2]}
    summary={**base['meta'],'fixed_radius_dry_time_s':fixed,'fixed_radius_dry_time_h':fixed/3600,'delta_geometry_s':float(t[-1]-fixed),'eta_geometry_percent':etaR,'eta_shape_percent':etaA,'E_shape_percent':max(map(abs,etaA.values())),'verification':verification,'excel_rows':len(t)-1,'excel_columns':22,'structural_empty_cells':int(np.isnan(arr[1:]).sum()),'table6':df.to_dict(orient='records'),'model_scope':'Q4 prescribed effective material-scalar diffusion, not identified solid-mass conservation'}
    # Serialize structural nulls as JSON null, never as invalid JSON NaN.
    clean=json.loads(json.dumps(summary,ensure_ascii=False).replace('NaN','null'))
    (res/'问题4_结果摘要.json').write_text(json.dumps(clean,indent=2,ensure_ascii=False,allow_nan=False),encoding='utf-8')
    return summary

def run_all(out,plots=True):
    out=out.resolve();res=out/'results';runpath=res/'q4_runs';runpath.mkdir(parents=True,exist_ok=True)
    total=time.perf_counter();runs={}
    def run(name,**kwargs):
        print('Computing',name,kwargs,flush=True);s=q.solve(**kwargs);validate_run(s)
        q.save_run(s,runpath/(name+'.npz'))
        print('Finished',name,s['meta']['dry_time_s'],s['meta']['runtime_s'],flush=True)
        return s
    N=320;base=run('main_320',N=N)
    fine=run('space_640',N=2*N)
    spatial_error=difference(base,fine);event_error=100*abs(base['time_s'][-1]/fine['time_s'][-1]-1)
    verification={'space_320_640':{'table_max_abs_kg_kg':spatial_error,'event_difference_s':float(base['time_s'][-1]-fine['time_s'][-1]),'event_relative_percent':event_error}}
    runs['main_320']=base;runs['space_640']=fine
    if spatial_error>5e-5 or event_error>.1:
        N=640;base=fine
        finer=run('space_1280',N=1280);runs['space_1280']=finer
        err=difference(base,finer);ev=100*abs(base['time_s'][-1]/finer['time_s'][-1]-1)
        verification['space_640_1280']={'table_max_abs_kg_kg':err,'event_difference_s':float(base['time_s'][-1]-finer['time_s'][-1]),'event_relative_percent':ev}
        assert err<=5e-5 and ev<=.1,'Spatial accuracy requires model review'
    runs['main']=base;q.save_run(base,runpath/'main.npz')
    ss=base['meta']['stable_air']
    cases=[('half_step',{'dt':.5}),('fixed',{'fixed':True}),('shape_0.1',{'a':.1}),('shape_0.2',{'a':.2}),('air_T_minus',{'Tshift':-ss['temperature_std_c']}),('air_T_plus',{'Tshift':ss['temperature_std_c']}),('air_C_minus',{'Cshift':-ss['moisture_std_kg_per_kg']}),('air_C_plus',{'Cshift':ss['moisture_std_kg_per_kg']}),('last_air',{'last_air':True})]
    # Independent frozen scenarios; two workers bound peak memory and CPU use.
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures=[(name,pool.submit(run,name,N=N,**kw)) for name,kw in cases]
        for name,future in futures:runs[name]=future.result()
    half=runs['half_step'];err=difference(base,half);ev=float(base['time_s'][-1]-half['time_s'][-1])
    verification['time_1_0.5']={'table_max_abs_kg_kg':err,'event_difference_s':ev,'event_relative_percent':100*abs(ev)/half['time_s'][-1]}
    assert abs(ev)<=1 and err<=5e-5,'Time refinement failed'
    assert runs['air_T_plus']['time_s'][-1]<runs['air_T_minus']['time_s'][-1]
    assert runs['air_C_plus']['time_s'][-1]>runs['air_C_minus']['time_s'][-1]
    summary=write_outputs(base,runs,out,verification)
    if plots:
        import 问题4_绘图 as plotting
        plotting.main(out)
    import 问题4_素材整理 as materials
    materials.main(out)
    manifest(out,summary)
    (res/'问题4_执行状态.json').write_text(json.dumps({'status':'computational_checks_passed','elapsed_s':time.perf_counter()-total,'figures_generated':plots},indent=2),encoding='utf-8')
    print(json.dumps({k:summary[k] for k in ['N','dry_time_s','dry_time_h','fixed_radius_dry_time_h','eta_geometry_percent','eta_shape_percent','E_shape_percent','verification']},ensure_ascii=False,indent=2),flush=True)

def manifest(out,summary):
    if (out/'results/复现清单.json').exists() and not (out/'results/复现清单_Q3历史.json').exists():
        shutil.copy2(out/'results/复现清单.json',out/'results/复现清单_Q3历史.json')
    inputs=[ROOT/'problem A/A题.pdf',q.old.AIR_PATH,q.old.RADIUS_PATH,q.old.TEMPLATE_DIR/'result4.xlsx',ROOT/'Q4/第四问建模思路与公式整理_最终推荐版.tex',ROOT/'题目分析报告.md',ROOT/'术语表格.md',ROOT/'问题4_求解.py',ROOT/'问题4_复现.py',ROOT/'问题4_绘图.py',ROOT/'问题4_素材整理.py',ROOT/'全程干燥求解.py',ROOT/'药材烘干求解.py',ROOT/'utils/plot_style.py']
    inputs += list((SKILL/'tools/figure/scripts').glob('*.py'))
    cmd=[sys.executable,'-X','utf8',str(SKILL/'references/roles/编程手/scripts/repro_manifest.py'),'--project-root',str(out),'--seed','0','--parameters',json.dumps({'question':4,'N':summary['N'],'dt_s':1.,'refinement_dt_s':.5,'shape_values':[0,.1,.2],'stable_start_s':9000,'stable_end_s':14400,'threshold':.15,'h':25.,'hm':8e-7}),'--command','python -X utf8 问题4_复现.py','--dependencies',json.dumps({'numba':'0.67.0','llvmlite':'0.49.0'}),'--overwrite']
    for f in inputs:cmd+=['--input',str(f)]
    for package in ['numpy','scipy','pandas','openpyxl','matplotlib']:cmd+=['--package',package]
    subprocess.run(cmd,check=True)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output-root',type=Path,default=ROOT);parser.add_argument('--no-plots',action='store_true')
    args=parser.parse_args();run_all(args.output_root,not args.no_plots)
