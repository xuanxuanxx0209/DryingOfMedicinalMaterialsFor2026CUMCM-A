from __future__ import annotations

import argparse
import hashlib
import json
import platform
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parent
AIR_PATH = PROJECT_ROOT / "problem A" / "附件" / "附件1.xlsx"
RADIUS_PATH = PROJECT_ROOT / "problem A" / "附件" / "附件2.xlsx"
TEMPLATE_DIR = PROJECT_ROOT / "problem A" / "附件" / "附件3"
RESULTS_DIR = PROJECT_ROOT / "results"


@dataclass(frozen=True)
class AirBoundary:
    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray

    def at(self, time_s: float) -> tuple[float, float]:
        temperature = float(
            np.interp(
                time_s,
                self.time_s,
                self.temperature_c,
                left=self.temperature_c[0],
                right=self.temperature_c[-1],
            )
        )
        moisture = float(
            np.interp(
                time_s,
                self.time_s,
                self.moisture,
                left=self.moisture[0],
                right=self.moisture[-1],
            )
        )
        return temperature, moisture


@dataclass(frozen=True)
class RadialGrid:
    radius_m: float
    length_m: float
    nodes_m: np.ndarray
    volumes_m3: np.ndarray
    internal_areas_m2: np.ndarray
    surface_area_m2: float
    spacing_m: float


@dataclass
class SimulationResult:
    time_s: np.ndarray
    radius_m: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    max_picard_iterations: int
    moisture_balance_relative_error: float
    heat_matrix_condition_number: float
    moisture_matrix_condition_number: float


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_air_boundary(path: Path = AIR_PATH) -> AirBoundary:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if len(rows) != 242:
        raise ValueError(f"附件1有效行数应为242，实际为{len(rows)}")
    header = tuple(rows[0])
    if header != ("时间", "温度", "水分浓度"):
        raise ValueError(f"附件1表头异常: {header}")
    data = np.asarray(rows[1:], dtype=float)
    if data.shape != (241, 3):
        raise ValueError(f"附件1数据形状异常: {data.shape}")
    if data[0, 0] != 0 or data[-1, 0] != 14400:
        raise ValueError("附件1时间范围不是0至14400秒")
    if not np.allclose(np.diff(data[:, 0]), 60.0):
        raise ValueError("附件1时间间隔不是60秒")
    if not np.all(np.isfinite(data)):
        raise ValueError("附件1含有非有限数")
    return AirBoundary(data[:, 0], data[:, 1], data[:, 2])


def make_grid(radius_m: float, intervals: int, length_m: float = 0.25) -> RadialGrid:
    if intervals < 2:
        raise ValueError("径向小区间数至少为2")
    nodes = np.linspace(0.0, radius_m, intervals + 1)
    spacing = radius_m / intervals
    faces = np.empty(intervals + 2, dtype=float)
    faces[0] = 0.0
    faces[-1] = radius_m
    faces[1:-1] = 0.5 * (nodes[:-1] + nodes[1:])
    volumes = np.pi * length_m * (faces[1:] ** 2 - faces[:-1] ** 2)
    internal_areas = 2.0 * np.pi * length_m * faces[1:-1]
    surface_area = 2.0 * np.pi * radius_m * length_m
    return RadialGrid(
        radius_m=radius_m,
        length_m=length_m,
        nodes_m=nodes,
        volumes_m3=volumes,
        internal_areas_m2=internal_areas,
        surface_area_m2=surface_area,
        spacing_m=spacing,
    )


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    n = diagonal.size
    c_prime = np.empty(n - 1, dtype=float)
    d_prime = np.empty(n, dtype=float)
    pivot = diagonal[0]
    if abs(pivot) < 1.0e-30:
        raise np.linalg.LinAlgError("三对角矩阵首个主元接近零")
    c_prime[0] = upper[0] / pivot
    d_prime[0] = right_hand_side[0] / pivot
    for i in range(1, n):
        pivot = diagonal[i] - lower[i - 1] * c_prime[i - 1]
        if abs(pivot) < 1.0e-30:
            raise np.linalg.LinAlgError(f"三对角矩阵第{i}个主元接近零")
        if i < n - 1:
            c_prime[i] = upper[i] / pivot
        d_prime[i] = (right_hand_side[i] - lower[i - 1] * d_prime[i - 1]) / pivot
    solution = np.empty(n, dtype=float)
    solution[-1] = d_prime[-1]
    for i in range(n - 2, -1, -1):
        solution[i] = d_prime[i] - c_prime[i] * solution[i + 1]
    return solution


def assemble_implicit_system(
    grid: RadialGrid,
    old_state: np.ndarray,
    storage_coefficient: np.ndarray,
    diffusion_coefficient: np.ndarray,
    time_step_s: float,
    boundary_coefficient: float,
    ambient_value: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    storage = storage_coefficient * grid.volumes_m3 / time_step_s
    face_diffusion = 0.5 * (
        diffusion_coefficient[:-1] + diffusion_coefficient[1:]
    )
    conductance = face_diffusion * grid.internal_areas_m2 / grid.spacing_m
    diagonal = storage.copy()
    diagonal[:-1] += conductance
    diagonal[1:] += conductance
    diagonal[-1] += boundary_coefficient * grid.surface_area_m2
    lower = -conductance.copy()
    upper = -conductance.copy()
    right_hand_side = storage * old_state
    right_hand_side[-1] += (
        boundary_coefficient * grid.surface_area_m2 * ambient_value
    )
    return lower, diagonal, upper, right_hand_side


def dense_from_tridiagonal(
    lower: np.ndarray, diagonal: np.ndarray, upper: np.ndarray
) -> np.ndarray:
    matrix = np.diag(diagonal)
    matrix += np.diag(lower, k=-1)
    matrix += np.diag(upper, k=1)
    return matrix


def q1_diffusivity(moisture: np.ndarray) -> np.ndarray:
    safe_moisture = np.maximum(moisture, 1.0e-8)
    return 7.0e-9 * np.exp(-0.89 / safe_moisture)


def simulate_q1(
    end_time_s: float,
    time_step_s: float = 1.0,
    intervals: int = 20,
    picard_tolerance: float = 1.0e-10,
    picard_max_iterations: int = 50,
) -> SimulationResult:
    steps_float = end_time_s / time_step_s
    steps = int(round(steps_float))
    if not np.isclose(steps, steps_float):
        raise ValueError("结束时刻必须是时间步长的整数倍")
    air = read_air_boundary()
    grid = make_grid(radius_m=0.02, intervals=intervals)
    node_count = intervals + 1
    time_axis = np.linspace(0.0, end_time_s, steps + 1)
    temperature_history = np.empty((steps + 1, node_count), dtype=float)
    moisture_history = np.empty((steps + 1, node_count), dtype=float)
    temperature = np.full(node_count, 28.0, dtype=float)
    moisture = np.full(node_count, 2.55, dtype=float)
    temperature_history[0] = temperature
    moisture_history[0] = moisture

    density = 820.0
    heat_capacity = 2600.0
    conductivity = 0.36
    heat_transfer_coefficient = 25.0
    mass_transfer_coefficient = 8.0e-7
    heat_storage = np.full(node_count, density * heat_capacity, dtype=float)
    heat_diffusion = np.full(node_count, conductivity, dtype=float)
    heat_lower, heat_diagonal, heat_upper, _ = assemble_implicit_system(
        grid,
        temperature,
        heat_storage,
        heat_diffusion,
        time_step_s,
        heat_transfer_coefficient,
        28.0,
    )
    heat_condition = float(
        np.linalg.cond(
            dense_from_tridiagonal(heat_lower, heat_diagonal, heat_upper)
        )
    )

    initial_inventory = float(np.dot(grid.volumes_m3, moisture))
    cumulative_surface_outflow = 0.0
    max_picard_iterations_used = 0
    moisture_condition = float("nan")

    for step in range(1, steps + 1):
        current_time = step * time_step_s
        ambient_temperature, ambient_moisture = air.at(current_time)
        _, _, _, heat_rhs = assemble_implicit_system(
            grid,
            temperature,
            heat_storage,
            heat_diffusion,
            time_step_s,
            heat_transfer_coefficient,
            ambient_temperature,
        )
        temperature = solve_tridiagonal(
            heat_lower, heat_diagonal, heat_upper, heat_rhs
        )

        moisture_guess = moisture.copy()
        converged = False
        for iteration in range(1, picard_max_iterations + 1):
            diffusivity = q1_diffusivity(moisture_guess)
            lower, diagonal, upper, rhs = assemble_implicit_system(
                grid,
                moisture,
                np.ones(node_count, dtype=float),
                diffusivity,
                time_step_s,
                mass_transfer_coefficient,
                ambient_moisture,
            )
            moisture_new = solve_tridiagonal(lower, diagonal, upper, rhs)
            update_norm = float(np.max(np.abs(moisture_new - moisture_guess)))
            moisture_guess = moisture_new
            if update_norm <= picard_tolerance:
                converged = True
                break
        if not converged:
            raise RuntimeError(
                f"问题一水分Picard迭代在t={current_time:g}秒未收敛"
            )
        if step == 1:
            moisture_condition = float(
                np.linalg.cond(dense_from_tridiagonal(lower, diagonal, upper))
            )
        max_picard_iterations_used = max(max_picard_iterations_used, iteration)
        moisture = moisture_guess
        cumulative_surface_outflow += (
            time_step_s
            * mass_transfer_coefficient
            * grid.surface_area_m2
            * (moisture[-1] - ambient_moisture)
        )
        temperature_history[step] = temperature
        moisture_history[step] = moisture

    final_inventory = float(np.dot(grid.volumes_m3, moisture))
    balance_residual = (
        final_inventory - initial_inventory + cumulative_surface_outflow
    )
    relative_balance_error = abs(balance_residual) / max(
        abs(initial_inventory - final_inventory), 1.0e-30
    )
    return SimulationResult(
        time_s=time_axis,
        radius_m=grid.nodes_m,
        temperature_c=temperature_history,
        moisture=moisture_history,
        max_picard_iterations=max_picard_iterations_used,
        moisture_balance_relative_error=float(relative_balance_error),
        heat_matrix_condition_number=heat_condition,
        moisture_matrix_condition_number=moisture_condition,
    )


def sample_field(
    result: SimulationResult,
    field: np.ndarray,
    sample_times_s: list[float],
    sample_radii_cm: list[float],
) -> np.ndarray:
    sampled = np.empty((len(sample_times_s), len(sample_radii_cm)), dtype=float)
    radius_cm = 100.0 * result.radius_m
    for i, sample_time in enumerate(sample_times_s):
        time_index = int(np.argmin(np.abs(result.time_s - sample_time)))
        if not np.isclose(result.time_s[time_index], sample_time):
            raise ValueError(f"计算结果不含要求时刻{sample_time}秒")
        sampled[i] = np.interp(sample_radii_cm, radius_cm, field[time_index])
    return sampled


def write_q1_workbook(result: SimulationResult, output_path: Path) -> None:
    template_path = TEMPLATE_DIR / "result1.xlsx"
    workbook = openpyxl.load_workbook(template_path)
    radii_cm = [round(0.1 * i, 1) for i in range(21)]
    output_times_s = np.arange(1.0, 1801.0, 1.0)
    time_indices = np.searchsorted(result.time_s, output_times_s)
    if not np.allclose(result.time_s[time_indices], output_times_s):
        raise ValueError("内部时间网格未包含result1要求的整数秒时刻")
    for sheet_name, field in (
        ("温度", result.temperature_c),
        ("水分浓度", result.moisture),
    ):
        worksheet = workbook[sheet_name]
        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)
        worksheet.cell(1, 1, "时间\\到药材中心的距离")
        for column, radius in enumerate(radii_cm, start=2):
            worksheet.cell(1, column, radius)
        for row, (time_value, time_index) in enumerate(
            zip(output_times_s, time_indices, strict=True), start=2
        ):
            worksheet.cell(row, 1, int(round(float(time_value))))
            values = np.interp(radii_cm, 100.0 * result.radius_m, field[time_index])
            for column, value in enumerate(values, start=2):
                worksheet.cell(row, column, round(float(value), 4))
        worksheet.freeze_panes = "B2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(".tmp.xlsx")
    workbook.save(temporary_path)
    workbook.close()
    temporary_path.replace(output_path)


def write_key_table_csv(
    path: Path,
    sample_times_s: list[float],
    sample_radii_cm: list[float],
    temperature: np.ndarray,
    moisture: np.ndarray,
) -> None:
    lines = ["变量,时间_s," + ",".join(f"r_{r:g}_cm" for r in sample_radii_cm)]
    for name, values in (("温度_C", temperature), ("水分浓度_kg_per_kg", moisture)):
        for time_value, row in zip(sample_times_s, values, strict=True):
            lines.append(
                f"{name},{time_value:g}," + ",".join(f"{value:.8f}" for value in row)
            )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def run_smoke() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    result = simulate_q1(end_time_s=120.0, time_step_s=1.0, intervals=20)
    ambient_120 = read_air_boundary().at(120.0)
    report = {
        "mode": "P1_q1_smoke",
        "input": {
            "path": str(AIR_PATH.relative_to(PROJECT_ROOT)),
            "sha256": sha256_file(AIR_PATH),
            "rows": 241,
            "first_time_s": 0,
            "last_time_s": 14400,
        },
        "parameters": {
            "end_time_s": 120.0,
            "time_step_s": 1.0,
            "radial_spacing_m": 0.001,
            "radial_intervals": 20,
            "initial_temperature_c": 28.0,
            "initial_moisture_kg_per_kg": 2.55,
            "diffusivity_formula": "7e-9*exp(-0.89/C)",
            "picard_tolerance": 1.0e-10,
            "picard_max_iterations": 50,
        },
        "boundary_at_120_s": {
            "temperature_c": ambient_120[0],
            "moisture_kg_per_kg": ambient_120[1],
        },
        "result_at_120_s": {
            "temperature_center_c": float(result.temperature_c[-1, 0]),
            "temperature_surface_c": float(result.temperature_c[-1, -1]),
            "moisture_center_kg_per_kg": float(result.moisture[-1, 0]),
            "moisture_surface_kg_per_kg": float(result.moisture[-1, -1]),
        },
        "checks": {
            "temperature_min_c": float(np.min(result.temperature_c)),
            "temperature_max_c": float(np.max(result.temperature_c)),
            "moisture_min_kg_per_kg": float(np.min(result.moisture)),
            "moisture_max_kg_per_kg": float(np.max(result.moisture)),
            "moisture_balance_relative_error": result.moisture_balance_relative_error,
            "heat_matrix_condition_number": result.heat_matrix_condition_number,
            "moisture_matrix_condition_number": result.moisture_matrix_condition_number,
            "max_picard_iterations": result.max_picard_iterations,
            "finite": bool(
                np.all(np.isfinite(result.temperature_c))
                and np.all(np.isfinite(result.moisture))
            ),
        },
    }
    output_path = RESULTS_DIR / "P1_q1_smoke.json"
    output_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


def run_q1() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    authoritative = simulate_q1(
        end_time_s=1800.0, time_step_s=0.25, intervals=160
    )
    time_refined = simulate_q1(
        end_time_s=1800.0, time_step_s=0.125, intervals=160
    )
    space_refined = simulate_q1(
        end_time_s=1800.0, time_step_s=0.25, intervals=320
    )
    sample_times = [100.0, 300.0, 600.0, 900.0, 1200.0, 1500.0, 1800.0]
    sample_radii = [0.0, 0.5, 1.0, 1.5, 2.0]
    authoritative_temperature = sample_field(
        authoritative, authoritative.temperature_c, sample_times, sample_radii
    )
    authoritative_moisture = sample_field(
        authoritative, authoritative.moisture, sample_times, sample_radii
    )
    time_refined_temperature = sample_field(
        time_refined, time_refined.temperature_c, sample_times, sample_radii
    )
    time_refined_moisture = sample_field(
        time_refined, time_refined.moisture, sample_times, sample_radii
    )
    space_refined_temperature = sample_field(
        space_refined, space_refined.temperature_c, sample_times, sample_radii
    )
    space_refined_moisture = sample_field(
        space_refined, space_refined.moisture, sample_times, sample_radii
    )
    time_temperature_difference = np.abs(
        authoritative_temperature - time_refined_temperature
    )
    time_moisture_difference = np.abs(
        authoritative_moisture - time_refined_moisture
    )
    space_temperature_difference = np.abs(
        authoritative_temperature - space_refined_temperature
    )
    space_moisture_difference = np.abs(
        authoritative_moisture - space_refined_moisture
    )

    result_workbook = RESULTS_DIR / "result1.xlsx"
    write_q1_workbook(authoritative, result_workbook)
    write_key_table_csv(
        RESULTS_DIR / "问题1_关键结果.csv",
        sample_times,
        sample_radii,
        authoritative_temperature,
        authoritative_moisture,
    )
    convergence_lines = [
        "检验,变量,最大绝对差,平均绝对差,权威空间步长_m,权威时间步长_s,加密空间步长_m,加密时间步长_s",
        f"时间步减半,温度_C,{time_temperature_difference.max():.12g},{time_temperature_difference.mean():.12g},0.000125,0.25,0.000125,0.125",
        f"时间步减半,水分浓度_kg_per_kg,{time_moisture_difference.max():.12g},{time_moisture_difference.mean():.12g},0.000125,0.25,0.000125,0.125",
        f"空间步减半,温度_C,{space_temperature_difference.max():.12g},{space_temperature_difference.mean():.12g},0.000125,0.25,0.0000625,0.25",
        f"空间步减半,水分浓度_kg_per_kg,{space_moisture_difference.max():.12g},{space_moisture_difference.mean():.12g},0.000125,0.25,0.0000625,0.25",
    ]
    (RESULTS_DIR / "问题1_网格收敛.csv").write_text(
        "\n".join(convergence_lines) + "\n", encoding="utf-8-sig"
    )

    summary = {
        "mode": "q1_full",
        "input_sha256": {str(AIR_PATH.relative_to(PROJECT_ROOT)): sha256_file(AIR_PATH)},
        "parameters": {
            "authoritative": {
                "radial_spacing_m": 0.000125,
                "time_step_s": 0.25,
            },
            "time_refined": {
                "radial_spacing_m": 0.000125,
                "time_step_s": 0.125,
            },
            "space_refined": {
                "radial_spacing_m": 0.0000625,
                "time_step_s": 0.25,
            },
            "diffusivity_formula": "7e-9*exp(-0.89/C)",
            "picard_tolerance": 1.0e-10,
        },
        "requested_times_s": sample_times,
        "requested_radii_cm": sample_radii,
        "temperature_c": np.round(authoritative_temperature, 8).tolist(),
        "moisture_kg_per_kg": np.round(authoritative_moisture, 8).tolist(),
        "verification": {
            "max_temperature_difference_time_refined_c": float(
                time_temperature_difference.max()
            ),
            "max_moisture_difference_time_refined_kg_per_kg": float(
                time_moisture_difference.max()
            ),
            "max_temperature_difference_space_refined_c": float(
                space_temperature_difference.max()
            ),
            "max_moisture_difference_space_refined_kg_per_kg": float(
                space_moisture_difference.max()
            ),
            "authoritative_moisture_balance_relative_error": authoritative.moisture_balance_relative_error,
            "time_refined_moisture_balance_relative_error": time_refined.moisture_balance_relative_error,
            "space_refined_moisture_balance_relative_error": space_refined.moisture_balance_relative_error,
            "authoritative_max_picard_iterations": authoritative.max_picard_iterations,
            "time_refined_max_picard_iterations": time_refined.max_picard_iterations,
            "space_refined_max_picard_iterations": space_refined.max_picard_iterations,
            "authoritative_heat_matrix_condition_number": authoritative.heat_matrix_condition_number,
            "authoritative_moisture_matrix_condition_number": authoritative.moisture_matrix_condition_number,
        },
        "outputs": {
            "workbook": str(result_workbook.relative_to(PROJECT_ROOT)),
            "key_table": "results/问题1_关键结果.csv",
            "convergence": "results/问题1_网格收敛.csv",
        },
        "environment": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "openpyxl": openpyxl.__version__,
            "platform": platform.platform(),
        },
    }
    (RESULTS_DIR / "问题1_结果摘要.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser(description="2026 CUMCM A题药材烘干数值求解")
    parser.add_argument(
        "--mode", choices=("smoke", "q1"), default="smoke", help="运行模式"
    )
    arguments = parser.parse_args()
    if arguments.mode == "smoke":
        run_smoke()
    else:
        run_q1()


if __name__ == "__main__":
    main()
