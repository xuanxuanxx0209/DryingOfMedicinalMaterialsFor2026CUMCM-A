from __future__ import annotations

import argparse
import csv
import hashlib
import json
import platform
import shutil
import time
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
from scipy.linalg import solve_banded

import 药材烘干求解 as core


PROJECT_ROOT = Path(__file__).resolve().parent
AIR_PATH = PROJECT_ROOT / "problem A" / "附件" / "附件1.xlsx"
RADIUS_PATH = PROJECT_ROOT / "problem A" / "附件" / "附件2.xlsx"
TEMPLATE_DIR = PROJECT_ROOT / "problem A" / "附件" / "附件3"
RESULTS_DIR = PROJECT_ROOT / "results"

HEAT_TRANSFER_COEFFICIENT = 25.0
MASS_TRANSFER_COEFFICIENT = 8.0e-7
INITIAL_TEMPERATURE_C = 28.0
INITIAL_MOISTURE = 2.55
DRYING_THRESHOLD = 0.15
STABLE_AIR_START_S = 9000.0
STABLE_AIR_END_S = 14400.0


@dataclass(frozen=True)
class RadiusHistory:
    time_s: np.ndarray
    radius_m: np.ndarray

    def at(self, time_s: float) -> float:
        return float(
            np.interp(
                time_s,
                self.time_s,
                self.radius_m,
                left=self.radius_m[0],
                right=self.radius_m[-1],
            )
        )


@dataclass
class VariableSimulation:
    time_s: np.ndarray
    radius_m: np.ndarray
    node_coordinate: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    max_coupling_iterations: int
    moisture_balance_relative_error: float
    dry_time_s: float | None
    coupling_iterations: np.ndarray
    max_temperature_linear_residual: float
    max_moisture_linear_residual: float


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def commit_temporary_file(temporary: Path, output_path: Path) -> None:
    try:
        temporary.replace(output_path)
    except PermissionError:
        shutil.copyfile(temporary, output_path)
        if sha256_file(temporary) != sha256_file(output_path):
            raise IOError(f"文件复制校验失败: {output_path}")
        temporary.unlink()


def read_radius_history(path: Path = RADIUS_PATH) -> RadiusHistory:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if len(rows) != 146:
        raise ValueError(f"附件2有效行数应为146，实际为{len(rows)}")
    if tuple(rows[0]) != ("时间", "半径"):
        raise ValueError(f"附件2表头异常: {tuple(rows[0])}")
    data = np.asarray(rows[1:], dtype=float)
    if data.shape != (145, 2):
        raise ValueError(f"附件2数据形状异常: {data.shape}")
    if data[0, 0] != 0 or data[-1, 0] != 259200:
        raise ValueError("附件2时间范围不是0至259200秒")
    if not np.allclose(np.diff(data[:, 0]), 1800.0):
        raise ValueError("附件2时间间隔不是1800秒")
    radius_m = data[:, 1] / 100.0
    if np.any(np.diff(radius_m) > 1.0e-12):
        raise ValueError("附件2半径存在增大记录")
    if not np.isclose(radius_m[0], 0.02) or not np.isclose(
        radius_m[-1], 0.01198
    ):
        raise ValueError("附件2首末半径异常")
    return RadiusHistory(data[:, 0], radius_m)


def q23_properties(
    moisture: np.ndarray, temperature_c: np.ndarray
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    safe_moisture = np.maximum(moisture, 1.0e-8)
    temperature_k = temperature_c + 273.15
    density = 650.0 + 128.0 * safe_moisture
    heat_capacity = 1450.0 + 2736.0 * safe_moisture / (safe_moisture + 1.0)
    conductivity = 0.21 + 0.38 * safe_moisture / (safe_moisture + 1.0)
    diffusivity = (
        2.4e-3
        * np.exp(-0.45 / safe_moisture)
        * np.exp(-3850.0 / temperature_k)
    )
    return density, heat_capacity, conductivity, diffusivity


def q4_properties(
    moisture: np.ndarray, temperature_c: np.ndarray
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    safe_moisture = np.maximum(moisture, 1.0e-8)
    temperature_k = temperature_c + 273.15
    density = 760.0 + 90.0 * safe_moisture
    heat_capacity = 1850.0 + 2150.0 * safe_moisture / (safe_moisture + 1.0)
    conductivity = 0.12 + 0.20 * safe_moisture / (safe_moisture + 1.0)
    diffusivity = (
        4.2e-4
        * np.exp(-0.30 / safe_moisture)
        * np.exp(-3850.0 / temperature_k)
    )
    return density, heat_capacity, conductivity, diffusivity


def stable_air_statistics() -> dict[str, float]:
    air = core.read_air_boundary()
    mask = (air.time_s >= STABLE_AIR_START_S) & (
        air.time_s <= STABLE_AIR_END_S
    )
    if int(np.sum(mask)) != 91:
        raise ValueError("附件1的9000至14400秒稳定段应含91个观测点")
    temperature = air.temperature_c[mask]
    moisture = air.moisture[mask]
    return {
        "sample_count": int(temperature.size),
        "temperature_mean_c": float(np.mean(temperature)),
        "temperature_std_c": float(np.std(temperature, ddof=0)),
        "temperature_min_c": float(np.min(temperature)),
        "temperature_max_c": float(np.max(temperature)),
        "moisture_mean_kg_per_kg": float(np.mean(moisture)),
        "moisture_std_kg_per_kg": float(np.std(moisture, ddof=0)),
        "moisture_min_kg_per_kg": float(np.min(moisture)),
        "moisture_max_kg_per_kg": float(np.max(moisture)),
        "last_temperature_c": float(air.temperature_c[-1]),
        "last_moisture_kg_per_kg": float(air.moisture[-1]),
    }


def solve_tridiagonal_fast(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    banded = np.zeros((3, diagonal.size), dtype=float)
    banded[0, 1:] = upper
    banded[1] = diagonal
    banded[2, :-1] = lower
    return solve_banded(
        (1, 1),
        banded,
        right_hand_side,
        overwrite_ab=True,
        overwrite_b=False,
        check_finite=False,
    )


def relative_tridiagonal_residual(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    solution: np.ndarray,
    right_hand_side: np.ndarray,
) -> float:
    residual = diagonal * solution - right_hand_side
    residual[1:] += lower * solution[:-1]
    residual[:-1] += upper * solution[1:]
    scale = max(float(np.max(np.abs(right_hand_side))), 1.0e-30)
    return float(np.max(np.abs(residual)) / scale)


def bdf2_effective_history(
    current_state: np.ndarray,
    previous_state: np.ndarray,
) -> np.ndarray:
    """Return the history state that lets the BE assembler realize BDF2.

    With an effective step of 2*dt/3, the standard three-level formula is
    (3*u[n+1] - 4*u[n] + u[n-1])/(2*dt). Material properties are evaluated
    at the current Picard iterate on time level n+1.
    """
    return (4.0 * current_state - previous_state) / 3.0


def simulate_variable_model(
    *,
    model: str,
    end_time_s: float,
    time_step_s: float,
    intervals: int,
    record_interval_s: float,
    stop_at_dryness: bool = False,
    coupling_temperature_tolerance: float = 1.0e-8,
    coupling_moisture_tolerance: float = 1.0e-10,
    coupling_max_iterations: int = 40,
    temporal_scheme: str = "backward_euler",
    heat_transfer_coefficient: float = HEAT_TRANSFER_COEFFICIENT,
    mass_transfer_coefficient: float = MASS_TRANSFER_COEFFICIENT,
    post_observation_temperature_c: float | None = None,
    post_observation_moisture: float | None = None,
) -> VariableSimulation:
    if model not in {"q23", "q4"}:
        raise ValueError("model必须为q23或q4")
    if temporal_scheme not in {"backward_euler", "bdf2"}:
        raise ValueError("temporal_scheme必须为backward_euler或bdf2")
    if heat_transfer_coefficient <= 0.0 or mass_transfer_coefficient <= 0.0:
        raise ValueError("表面对流系数必须为正数")
    if (post_observation_temperature_c is None) != (
        post_observation_moisture is None
    ):
        raise ValueError("长期空气温度和水分边界必须同时给出或同时省略")
    if post_observation_moisture is not None and post_observation_moisture < 0.0:
        raise ValueError("长期空气水分浓度不得为负")
    step_count_float = end_time_s / time_step_s
    step_count = int(round(step_count_float))
    if not np.isclose(step_count, step_count_float):
        raise ValueError("结束时刻必须是时间步长的整数倍")
    record_stride_float = record_interval_s / time_step_s
    record_stride = int(round(record_stride_float))
    if not np.isclose(record_stride, record_stride_float):
        raise ValueError("记录间隔必须是时间步长的整数倍")

    air = core.read_air_boundary()
    radius_history = read_radius_history() if model == "q4" else None
    property_function = q23_properties if model == "q23" else q4_properties
    node_coordinate = np.linspace(0.0, 1.0, intervals + 1)
    temperature = np.full(intervals + 1, INITIAL_TEMPERATURE_C, dtype=float)
    moisture = np.full(intervals + 1, INITIAL_MOISTURE, dtype=float)

    recorded_time = [0.0]
    initial_radius = 0.02
    recorded_radius = [initial_radius]
    recorded_temperature = [temperature.copy()]
    recorded_moisture = [moisture.copy()]
    recorded_iterations = [0]
    normalized_weights = np.empty(intervals + 1, dtype=float)
    normalized_faces = np.empty(intervals + 2, dtype=float)
    normalized_faces[0] = 0.0
    normalized_faces[-1] = 1.0
    normalized_faces[1:-1] = 0.5 * (
        node_coordinate[:-1] + node_coordinate[1:]
    )
    normalized_weights[:] = (
        normalized_faces[1:] ** 2 - normalized_faces[:-1] ** 2
    )
    if not np.isclose(np.sum(normalized_weights), 1.0):
        raise RuntimeError("归一化控制体权重之和不等于1")

    previous_temperature = temperature.copy()
    previous_moisture = moisture.copy()
    max_iterations_used = 0
    max_balance_error = 0.0
    max_temperature_linear_residual = 0.0
    max_moisture_linear_residual = 0.0
    dry_time_s: float | None = None

    for step in range(1, step_count + 1):
        current_time = step * time_step_s
        radius_m = (
            0.02 if radius_history is None else radius_history.at(current_time)
        )
        grid = core.make_grid(radius_m=radius_m, intervals=intervals)
        ambient_temperature, ambient_moisture = air.at(current_time)
        if (
            current_time > float(air.time_s[-1])
            and post_observation_temperature_c is not None
            and post_observation_moisture is not None
        ):
            ambient_temperature = post_observation_temperature_c
            ambient_moisture = post_observation_moisture
        temperature_guess = temperature.copy()
        moisture_guess = moisture.copy()
        converged = False

        for iteration in range(1, coupling_max_iterations + 1):
            use_bdf2 = temporal_scheme == "bdf2" and step > 1
            effective_step_s = 2.0 * time_step_s / 3.0 if use_bdf2 else time_step_s
            temperature_history_state = (
                bdf2_effective_history(temperature, previous_temperature)
                if use_bdf2
                else temperature
            )
            moisture_history_state = (
                bdf2_effective_history(moisture, previous_moisture)
                if use_bdf2
                else moisture
            )
            density, heat_capacity, conductivity, _ = property_function(
                moisture_guess, temperature_guess
            )
            lower_t, diagonal_t, upper_t, rhs_t = core.assemble_implicit_system(
                grid,
                temperature_history_state,
                density * heat_capacity,
                conductivity,
                effective_step_s,
                heat_transfer_coefficient,
                ambient_temperature,
            )
            temperature_new = solve_tridiagonal_fast(
                lower_t, diagonal_t, upper_t, rhs_t
            )
            max_temperature_linear_residual = max(
                max_temperature_linear_residual,
                relative_tridiagonal_residual(
                    lower_t, diagonal_t, upper_t, temperature_new, rhs_t
                ),
            )
            _, _, _, diffusivity = property_function(
                moisture_guess, temperature_new
            )
            lower_c, diagonal_c, upper_c, rhs_c = core.assemble_implicit_system(
                grid,
                moisture_history_state,
                np.ones(intervals + 1, dtype=float),
                diffusivity,
                effective_step_s,
                mass_transfer_coefficient,
                ambient_moisture,
            )
            moisture_new = solve_tridiagonal_fast(
                lower_c, diagonal_c, upper_c, rhs_c
            )
            max_moisture_linear_residual = max(
                max_moisture_linear_residual,
                relative_tridiagonal_residual(
                    lower_c, diagonal_c, upper_c, moisture_new, rhs_c
                ),
            )
            temperature_error = float(
                np.max(np.abs(temperature_new - temperature_guess))
            )
            moisture_error = float(np.max(np.abs(moisture_new - moisture_guess)))
            temperature_guess = temperature_new
            moisture_guess = moisture_new
            if (
                temperature_error <= coupling_temperature_tolerance
                and moisture_error <= coupling_moisture_tolerance
            ):
                converged = True
                break
        if not converged:
            raise RuntimeError(
                f"{model}耦合迭代在t={current_time:g}秒未收敛，"
                f"温度误差{temperature_error:.3e}，含水率误差{moisture_error:.3e}"
            )

        mean_new = float(np.dot(normalized_weights, moisture_guess))
        mean_current = float(np.dot(normalized_weights, moisture))
        if temporal_scheme == "bdf2" and step > 1:
            mean_previous = float(np.dot(normalized_weights, previous_moisture))
            inventory_change = 1.5 * mean_new - 2.0 * mean_current + 0.5 * mean_previous
        else:
            inventory_change = mean_new - mean_current
        boundary_loss = (
            time_step_s
            * 2.0
            * mass_transfer_coefficient
            / radius_m
            * (moisture_guess[-1] - ambient_moisture)
        )
        balance_scale = max(abs(inventory_change), abs(boundary_loss), 1.0e-30)
        max_balance_error = max(
            max_balance_error,
            abs(inventory_change + boundary_loss) / balance_scale,
        )

        previous_temperature, temperature = temperature, temperature_guess
        previous_moisture, moisture = moisture, moisture_guess
        max_iterations_used = max(max_iterations_used, iteration)

        if not np.all(np.isfinite(temperature)) or not np.all(np.isfinite(moisture)):
            raise RuntimeError(f"{model}在t={current_time:g}秒出现非有限状态")
        if float(np.min(moisture)) < -1.0e-12:
            raise RuntimeError(f"{model}在t={current_time:g}秒出现负含水率")

        if step % record_stride == 0:
            recorded_time.append(current_time)
            recorded_radius.append(radius_m)
            recorded_temperature.append(temperature.copy())
            recorded_moisture.append(moisture.copy())
            recorded_iterations.append(iteration)

        if stop_at_dryness and float(np.max(moisture)) < DRYING_THRESHOLD:
            dry_time_s = current_time
            if recorded_time[-1] != current_time:
                recorded_time.append(current_time)
                recorded_radius.append(radius_m)
                recorded_temperature.append(temperature.copy())
                recorded_moisture.append(moisture.copy())
                recorded_iterations.append(iteration)
            break

    if stop_at_dryness and dry_time_s is None:
        raise RuntimeError(f"{model}在{end_time_s / 3600:g}小时内未达到干燥阈值")

    return VariableSimulation(
        time_s=np.asarray(recorded_time, dtype=float),
        radius_m=np.asarray(recorded_radius, dtype=float),
        node_coordinate=node_coordinate,
        temperature_c=np.asarray(recorded_temperature, dtype=float),
        moisture=np.asarray(recorded_moisture, dtype=float),
        max_coupling_iterations=max_iterations_used,
        moisture_balance_relative_error=float(max_balance_error),
        dry_time_s=dry_time_s,
        coupling_iterations=np.asarray(recorded_iterations, dtype=int),
        max_temperature_linear_residual=float(max_temperature_linear_residual),
        max_moisture_linear_residual=float(max_moisture_linear_residual),
    )


def interpolate_at_physical_radii(
    simulation: VariableSimulation,
    field: np.ndarray,
    row_index: int,
    radii_cm: list[float],
) -> np.ndarray:
    current_radius_cm = 100.0 * simulation.radius_m[row_index]
    physical_nodes_cm = simulation.node_coordinate * current_radius_cm
    return np.interp(radii_cm, physical_nodes_cm, field[row_index])


def write_result2(simulation: VariableSimulation, output_path: Path) -> None:
    if not np.allclose(np.diff(simulation.time_s), 1.0):
        raise ValueError("result2要求每1秒记录一次")
    template = openpyxl.load_workbook(TEMPLATE_DIR / "result2.xlsx")
    radii_cm = [round(0.1 * index, 1) for index in range(21)]
    for sheet_name, field in (
        ("温度", simulation.temperature_c),
        ("水分浓度", simulation.moisture),
    ):
        worksheet = template[sheet_name]
        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)
        worksheet.cell(1, 1, "时间\\到药材中心的距离")
        for column, radius_cm in enumerate(radii_cm, start=2):
            worksheet.cell(1, column, radius_cm)
        for row, time_s in enumerate(simulation.time_s[1:], start=2):
            source_index = row - 1
            values = interpolate_at_physical_radii(
                simulation, field, source_index, radii_cm
            )
            worksheet.cell(row, 1, int(round(float(time_s))))
            for column, value in enumerate(values, start=2):
                cell = worksheet.cell(row, column, round(float(value), 4))
                cell.number_format = "0.0000"
        worksheet.freeze_panes = "B2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".tmp.xlsx")
    template.save(temporary)
    template.close()
    commit_temporary_file(temporary, output_path)


def write_result3(simulation: VariableSimulation, output_path: Path) -> None:
    template = openpyxl.load_workbook(TEMPLATE_DIR / "result3.xlsx")
    worksheet = template.active
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    radii_cm = [round(0.1 * index, 1) for index in range(21)]
    worksheet.cell(1, 1, "时间\\到药材中心的距离")
    for column, radius_cm in enumerate(radii_cm, start=2):
        worksheet.cell(1, column, radius_cm)
    for row, time_s in enumerate(simulation.time_s[1:], start=2):
        source_index = row - 1
        values = interpolate_at_physical_radii(
            simulation, simulation.moisture, source_index, radii_cm
        )
        worksheet.cell(row, 1, int(round(float(time_s))))
        for column, value in enumerate(values, start=2):
            cell = worksheet.cell(row, column, round(float(value), 4))
            cell.number_format = "0.0000"
    worksheet.freeze_panes = "B2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".tmp.xlsx")
    template.save(temporary)
    template.close()
    commit_temporary_file(temporary, output_path)


def write_result4(simulation: VariableSimulation, output_path: Path) -> None:
    template = openpyxl.load_workbook(TEMPLATE_DIR / "result4.xlsx")
    worksheet = template.active
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    fixed_radii_cm = [round(0.1 * index, 1) for index in range(20)]
    worksheet.cell(1, 1, "时间\\到药材中心的距离")
    for column, radius_cm in enumerate(fixed_radii_cm, start=2):
        worksheet.cell(1, column, radius_cm)
    surface_column = len(fixed_radii_cm) + 2
    worksheet.cell(1, surface_column, "药材表面")

    for row, time_s in enumerate(simulation.time_s[1:], start=2):
        source_index = row - 1
        current_radius_cm = 100.0 * simulation.radius_m[source_index]
        physical_nodes_cm = simulation.node_coordinate * current_radius_cm
        worksheet.cell(row, 1, int(round(float(time_s))))
        for column, radius_cm in enumerate(fixed_radii_cm, start=2):
            if radius_cm < current_radius_cm - 1.0e-10:
                value = np.interp(
                    radius_cm,
                    physical_nodes_cm,
                    simulation.moisture[source_index],
                )
                worksheet.cell(row, column, round(float(value), 4))
            else:
                worksheet.cell(row, column, None)
        worksheet.cell(
            row,
            surface_column,
            round(float(simulation.moisture[source_index, -1]), 4),
        )
    worksheet.freeze_panes = "B2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".tmp.xlsx")
    template.save(temporary)
    template.close()
    commit_temporary_file(temporary, output_path)


def requested_table(
    simulation: VariableSimulation,
    field: np.ndarray,
    times_s: list[float],
    radii_cm: list[float],
) -> list[list[float]]:
    rows: list[list[float]] = []
    for time_s in times_s:
        index = int(np.argmin(np.abs(simulation.time_s - time_s)))
        if not np.isclose(simulation.time_s[index], time_s):
            raise ValueError(f"记录结果不含时刻{time_s}秒")
        rows.append(
            interpolate_at_physical_radii(
                simulation, field, index, radii_cm
            ).tolist()
        )
    return rows


def write_table_csv(
    output_path: Path,
    times_h: list[float],
    radii_cm: list[float],
    values: np.ndarray,
    value_format: str = ".8f",
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间_h", *[f"{radius:g}_cm" for radius in radii_cm]])
        for time_h, row in zip(times_h, values, strict=True):
            writer.writerow([f"{time_h:.1f}", *[format(float(value), value_format) for value in row]])


def final_system_condition_numbers(
    simulation: VariableSimulation,
    time_step_s: float,
    heat_transfer_coefficient: float,
    mass_transfer_coefficient: float,
    post_observation_temperature_c: float | None = None,
    post_observation_moisture: float | None = None,
) -> dict[str, float]:
    intervals = simulation.node_coordinate.size - 1
    radius_m = float(simulation.radius_m[-1])
    grid = core.make_grid(radius_m=radius_m, intervals=intervals)
    air = core.read_air_boundary()
    ambient_temperature, ambient_moisture = air.at(float(simulation.time_s[-1]))
    if (
        float(simulation.time_s[-1]) > float(air.time_s[-1])
        and post_observation_temperature_c is not None
        and post_observation_moisture is not None
    ):
        ambient_temperature = post_observation_temperature_c
        ambient_moisture = post_observation_moisture
    density, heat_capacity, conductivity, diffusivity = q23_properties(
        simulation.moisture[-1], simulation.temperature_c[-1]
    )
    effective_step_s = 2.0 * time_step_s / 3.0
    lower_t, diagonal_t, upper_t, _ = core.assemble_implicit_system(
        grid,
        simulation.temperature_c[-1],
        density * heat_capacity,
        conductivity,
        effective_step_s,
        heat_transfer_coefficient,
        ambient_temperature,
    )
    lower_c, diagonal_c, upper_c, _ = core.assemble_implicit_system(
        grid,
        simulation.moisture[-1],
        np.ones(intervals + 1, dtype=float),
        diffusivity,
        effective_step_s,
        mass_transfer_coefficient,
        ambient_moisture,
    )
    return {
        "temperature_final": float(
            np.linalg.cond(core.dense_from_tridiagonal(lower_t, diagonal_t, upper_t))
        ),
        "moisture_final": float(
            np.linalg.cond(core.dense_from_tridiagonal(lower_c, diagonal_c, upper_c))
        ),
    }


def run_smoke() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    q2 = simulate_variable_model(
        model="q23",
        end_time_s=600.0,
        time_step_s=1.0,
        intervals=20,
        record_interval_s=60.0,
        temporal_scheme="bdf2",
    )
    q4 = simulate_variable_model(
        model="q4",
        end_time_s=3600.0,
        time_step_s=10.0,
        intervals=20,
        record_interval_s=60.0,
    )
    q2_rho, q2_cp, q2_k, q2_d = q23_properties(
        q2.moisture[-1], q2.temperature_c[-1]
    )
    q4_rho, q4_cp, q4_k, q4_d = q4_properties(
        q4.moisture[-1], q4.temperature_c[-1]
    )
    report = {
        "mode": "P1_q2_q4_smoke",
        "inputs": {
            str(AIR_PATH.relative_to(PROJECT_ROOT)): sha256_file(AIR_PATH),
            str(RADIUS_PATH.relative_to(PROJECT_ROOT)): sha256_file(RADIUS_PATH),
        },
        "q2": {
            "time_step_s": 1.0,
            "radial_intervals": 20,
            "end_time_s": 600.0,
            "temperature_range_c": [
                float(np.min(q2.temperature_c)),
                float(np.max(q2.temperature_c)),
            ],
            "moisture_range_kg_per_kg": [
                float(np.min(q2.moisture)),
                float(np.max(q2.moisture)),
            ],
            "property_ranges": {
                "density_kg_per_m3": [float(q2_rho.min()), float(q2_rho.max())],
                "heat_capacity_j_per_kg_k": [float(q2_cp.min()), float(q2_cp.max())],
                "conductivity_w_per_m_k": [float(q2_k.min()), float(q2_k.max())],
                "diffusivity_m2_per_s": [float(q2_d.min()), float(q2_d.max())],
            },
            "max_coupling_iterations": q2.max_coupling_iterations,
            "moisture_balance_relative_error": q2.moisture_balance_relative_error,
            "max_temperature_linear_residual": q2.max_temperature_linear_residual,
            "max_moisture_linear_residual": q2.max_moisture_linear_residual,
        },
        "q4": {
            "time_step_s": 10.0,
            "radial_intervals": 20,
            "end_time_s": 3600.0,
            "radius_start_m": float(q4.radius_m[0]),
            "radius_end_m": float(q4.radius_m[-1]),
            "temperature_range_c": [
                float(np.min(q4.temperature_c)),
                float(np.max(q4.temperature_c)),
            ],
            "moisture_range_kg_per_kg": [
                float(np.min(q4.moisture)),
                float(np.max(q4.moisture)),
            ],
            "property_ranges": {
                "density_kg_per_m3": [float(q4_rho.min()), float(q4_rho.max())],
                "heat_capacity_j_per_kg_k": [float(q4_cp.min()), float(q4_cp.max())],
                "conductivity_w_per_m_k": [float(q4_k.min()), float(q4_k.max())],
                "diffusivity_m2_per_s": [float(q4_d.min()), float(q4_d.max())],
            },
            "max_coupling_iterations": q4.max_coupling_iterations,
            "moisture_balance_relative_error": q4.moisture_balance_relative_error,
        },
    }
    output = RESULTS_DIR / "P1_q2_q4_smoke.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


def run_q2() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    authoritative = simulate_variable_model(
        model="q23",
        end_time_s=10800.0,
        time_step_s=1.0,
        intervals=160,
        record_interval_s=1.0,
        temporal_scheme="bdf2",
    )
    time_refined = simulate_variable_model(
        model="q23",
        end_time_s=10800.0,
        time_step_s=0.5,
        intervals=160,
        record_interval_s=1800.0,
        temporal_scheme="bdf2",
    )
    space_refined = simulate_variable_model(
        model="q23",
        end_time_s=10800.0,
        time_step_s=1.0,
        intervals=320,
        record_interval_s=1800.0,
        temporal_scheme="bdf2",
    )
    write_result2(authoritative, RESULTS_DIR / "result2.xlsx")

    times_s = [1800.0 * index for index in range(1, 7)]
    radii_cm = [0.0, 0.5, 1.0, 1.5, 2.0]
    temperature = np.asarray(
        requested_table(
            authoritative, authoritative.temperature_c, times_s, radii_cm
        )
    )
    moisture = np.asarray(
        requested_table(authoritative, authoritative.moisture, times_s, radii_cm)
    )
    time_temperature = np.asarray(
        requested_table(
            time_refined, time_refined.temperature_c, times_s, radii_cm
        )
    )
    time_moisture = np.asarray(
        requested_table(time_refined, time_refined.moisture, times_s, radii_cm)
    )
    space_temperature = np.asarray(
        requested_table(
            space_refined, space_refined.temperature_c, times_s, radii_cm
        )
    )
    space_moisture = np.asarray(
        requested_table(space_refined, space_refined.moisture, times_s, radii_cm)
    )
    times_h = [time_s / 3600.0 for time_s in times_s]
    write_table_csv(
        RESULTS_DIR / "问题2_表3温度.csv", times_h, radii_cm, temperature
    )
    write_table_csv(
        RESULTS_DIR / "问题2_表4含水率.csv", times_h, radii_cm, moisture
    )

    with (RESULTS_DIR / "问题2_迭代次数.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间_s", "Picard迭代次数"])
        writer.writerows(
            zip(
                authoritative.time_s.astype(int).tolist(),
                authoritative.coupling_iterations.tolist(),
                strict=True,
            )
        )

    property_density, property_cp, property_k, property_d = q23_properties(
        authoritative.moisture, authoritative.temperature_c
    )
    property_rows = []
    for time_s in times_s:
        row_index = int(np.argmin(np.abs(authoritative.time_s - time_s)))
        property_rows.append(
            [
                time_s / 3600.0,
                float(np.min(property_density[row_index])),
                float(np.max(property_density[row_index])),
                float(np.min(property_cp[row_index])),
                float(np.max(property_cp[row_index])),
                float(np.min(property_k[row_index])),
                float(np.max(property_k[row_index])),
                float(np.min(property_d[row_index])),
                float(np.max(property_d[row_index])),
            ]
        )
    with (RESULTS_DIR / "问题2_物性范围.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "时间_h",
                "密度最小_kg_m3",
                "密度最大_kg_m3",
                "比热最小_J_kg_K",
                "比热最大_J_kg_K",
                "导热系数最小_W_m_K",
                "导热系数最大_W_m_K",
                "扩散系数最小_m2_s",
                "扩散系数最大_m2_s",
            ]
        )
        writer.writerows(property_rows)

    sensitivity_cases = {
        "换热系数_0.8倍": (0.8 * HEAT_TRANSFER_COEFFICIENT, MASS_TRANSFER_COEFFICIENT),
        "换热系数_1.2倍": (1.2 * HEAT_TRANSFER_COEFFICIENT, MASS_TRANSFER_COEFFICIENT),
        "传质系数_0.8倍": (HEAT_TRANSFER_COEFFICIENT, 0.8 * MASS_TRANSFER_COEFFICIENT),
        "传质系数_1.2倍": (HEAT_TRANSFER_COEFFICIENT, 1.2 * MASS_TRANSFER_COEFFICIENT),
    }
    sensitivity = []
    for case_name, (heat_coefficient, mass_coefficient) in sensitivity_cases.items():
        case_simulation = simulate_variable_model(
            model="q23",
            end_time_s=10800.0,
            time_step_s=1.0,
            intervals=160,
            record_interval_s=1800.0,
            temporal_scheme="bdf2",
            heat_transfer_coefficient=heat_coefficient,
            mass_transfer_coefficient=mass_coefficient,
        )
        case_temperature = np.asarray(
            requested_table(
                case_simulation,
                case_simulation.temperature_c,
                times_s,
                radii_cm,
            )
        )
        case_moisture = np.asarray(
            requested_table(
                case_simulation,
                case_simulation.moisture,
                times_s,
                radii_cm,
            )
        )
        sensitivity.append(
            {
                "case": case_name,
                "heat_transfer_coefficient_w_per_m2_k": heat_coefficient,
                "mass_transfer_coefficient_m_per_s": mass_coefficient,
                "max_temperature_change_c": float(
                    np.max(np.abs(case_temperature - temperature))
                ),
                "max_moisture_change_kg_per_kg": float(
                    np.max(np.abs(case_moisture - moisture))
                ),
                "temperature_center_3h_change_c": float(
                    case_temperature[-1, 0] - temperature[-1, 0]
                ),
                "temperature_surface_3h_change_c": float(
                    case_temperature[-1, -1] - temperature[-1, -1]
                ),
                "moisture_center_3h_change_kg_per_kg": float(
                    case_moisture[-1, 0] - moisture[-1, 0]
                ),
                "moisture_surface_3h_change_kg_per_kg": float(
                    case_moisture[-1, -1] - moisture[-1, -1]
                ),
            }
        )
    with (RESULTS_DIR / "问题2_敏感性.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        fieldnames = list(sensitivity[0].keys())
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(sensitivity)

    condition_numbers = final_system_condition_numbers(
        authoritative,
        time_step_s=1.0,
        heat_transfer_coefficient=HEAT_TRANSFER_COEFFICIENT,
        mass_transfer_coefficient=MASS_TRANSFER_COEFFICIENT,
    )
    summary = {
        "mode": "q2_full",
        "inputs": {
            str(AIR_PATH.relative_to(PROJECT_ROOT)): sha256_file(AIR_PATH)
        },
        "parameters": {
            "temporal_scheme": "backward_euler_start_then_bdf2",
            "authoritative": {"intervals": 160, "time_step_s": 1.0},
            "time_refined": {"intervals": 160, "time_step_s": 0.5},
            "space_refined": {"intervals": 320, "time_step_s": 1.0},
            "heat_transfer_coefficient_w_per_m2_k": HEAT_TRANSFER_COEFFICIENT,
            "mass_transfer_coefficient_m_per_s": MASS_TRANSFER_COEFFICIENT,
            "coupling_temperature_tolerance_c": 1.0e-8,
            "coupling_moisture_tolerance_kg_per_kg": 1.0e-10,
            "coupling_max_iterations": 40,
        },
        "requested_times_s": times_s,
        "requested_radii_cm": radii_cm,
        "temperature_c": np.round(temperature, 8).tolist(),
        "moisture_kg_per_kg": np.round(moisture, 8).tolist(),
        "property_ranges": {
            "density_kg_per_m3": [
                float(np.min(property_density)),
                float(np.max(property_density)),
            ],
            "heat_capacity_j_per_kg_k": [
                float(np.min(property_cp)),
                float(np.max(property_cp)),
            ],
            "conductivity_w_per_m_k": [
                float(np.min(property_k)),
                float(np.max(property_k)),
            ],
            "diffusivity_m2_per_s": [
                float(np.min(property_d)),
                float(np.max(property_d)),
            ],
        },
        "sensitivity": sensitivity,
        "verification": {
            "max_temperature_difference_time_refined_c": float(
                np.max(np.abs(temperature - time_temperature))
            ),
            "max_moisture_difference_time_refined_kg_per_kg": float(
                np.max(np.abs(moisture - time_moisture))
            ),
            "max_temperature_difference_space_refined_c": float(
                np.max(np.abs(temperature - space_temperature))
            ),
            "max_moisture_difference_space_refined_kg_per_kg": float(
                np.max(np.abs(moisture - space_moisture))
            ),
            "rounded_temperature_changes_time_refined": int(
                np.sum(np.round(temperature, 4) != np.round(time_temperature, 4))
            ),
            "rounded_moisture_changes_time_refined": int(
                np.sum(np.round(moisture, 4) != np.round(time_moisture, 4))
            ),
            "rounded_temperature_changes_space_refined": int(
                np.sum(np.round(temperature, 4) != np.round(space_temperature, 4))
            ),
            "rounded_moisture_changes_space_refined": int(
                np.sum(np.round(moisture, 4) != np.round(space_moisture, 4))
            ),
            "moisture_balance_relative_error": authoritative.moisture_balance_relative_error,
            "max_coupling_iterations": authoritative.max_coupling_iterations,
            "max_temperature_linear_residual": authoritative.max_temperature_linear_residual,
            "max_moisture_linear_residual": authoritative.max_moisture_linear_residual,
            "final_system_condition_numbers": condition_numbers,
        },
        "environment": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "openpyxl": openpyxl.__version__,
            "platform": platform.platform(),
        },
    }
    (RESULTS_DIR / "问题2_结果摘要.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


def q3_periodic_table(
    simulation: VariableSimulation,
    radii_cm: list[float],
) -> tuple[list[float], np.ndarray]:
    if simulation.dry_time_s is None:
        raise RuntimeError("问题三未返回烘干结束时刻")
    periodic_times = list(
        np.arange(6.0 * 3600.0, simulation.dry_time_s, 6.0 * 3600.0)
    )
    table_times = periodic_times + [simulation.dry_time_s]
    table = np.asarray(
        requested_table(simulation, simulation.moisture, table_times, radii_cm)
    )
    return table_times, table


def q3_common_table_difference(
    first: VariableSimulation,
    second: VariableSimulation,
    radii_cm: list[float],
) -> float:
    if first.dry_time_s is None or second.dry_time_s is None:
        raise RuntimeError("网格比较缺少烘干结束时刻")
    common_end = min(first.dry_time_s, second.dry_time_s)
    periodic_times = list(np.arange(6.0 * 3600.0, common_end, 6.0 * 3600.0))
    first_table = np.asarray(
        requested_table(first, first.moisture, periodic_times, radii_cm)
    )
    second_table = np.asarray(
        requested_table(second, second.moisture, periodic_times, radii_cm)
    )
    periodic_difference = float(np.max(np.abs(first_table - second_table)))
    first_end = interpolate_at_physical_radii(
        first, first.moisture, -1, radii_cm
    )
    second_end = interpolate_at_physical_radii(
        second, second.moisture, -1, radii_cm
    )
    return max(periodic_difference, float(np.max(np.abs(first_end - second_end))))


def write_q3_process_csv(
    simulation: VariableSimulation,
    output_path: Path,
) -> None:
    radii_cm = [0.0, 0.5, 1.0, 1.5, 2.0]
    with output_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "time_s",
                "time_h",
                *[f"moisture_{radius:g}_cm" for radius in radii_cm],
                "maximum_moisture",
                "maximum_radius_cm",
            ]
        )
        for index, time_s in enumerate(simulation.time_s):
            values = interpolate_at_physical_radii(
                simulation, simulation.moisture, index, radii_cm
            )
            maximum_index = int(np.argmax(simulation.moisture[index]))
            writer.writerow(
                [
                    f"{time_s:.6f}",
                    f"{time_s / 3600.0:.10f}",
                    *[f"{value:.10f}" for value in values],
                    f"{simulation.moisture[index, maximum_index]:.10f}",
                    f"{100.0 * simulation.radius_m[index] * simulation.node_coordinate[maximum_index]:.8f}",
                ]
            )


def run_q3_smoke() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    stable = stable_air_statistics()
    simulation = simulate_variable_model(
        model="q23",
        end_time_s=18000.0,
        time_step_s=10.0,
        intervals=20,
        record_interval_s=60.0,
        temporal_scheme="bdf2",
        post_observation_temperature_c=stable["temperature_mean_c"],
        post_observation_moisture=stable["moisture_mean_kg_per_kg"],
    )
    maximum_series = np.max(simulation.moisture, axis=1)
    maximum_increase = float(np.max(np.diff(maximum_series)))
    if maximum_increase > 1.0e-8:
        raise RuntimeError("问题三冒烟测试的最大含水率出现异常回升")
    if simulation.moisture_balance_relative_error > 1.0e-8:
        raise RuntimeError("问题三冒烟测试的水分平衡残差超过合同阈值")
    if simulation.max_temperature_linear_residual > 1.0e-12:
        raise RuntimeError("问题三冒烟测试的温度代数残差超过合同阈值")
    if simulation.max_moisture_linear_residual > 1.0e-12:
        raise RuntimeError("问题三冒烟测试的含水率代数残差超过合同阈值")
    report = {
        "mode": "P1_q3_smoke",
        "inputs": {
            str(AIR_PATH.relative_to(PROJECT_ROOT)): sha256_file(AIR_PATH),
        },
        "model": "q23_fixed_radius_stable_mean_extension",
        "temporal_scheme": "backward_euler_start_then_bdf2",
        "time_step_s": 10.0,
        "radial_intervals": 20,
        "end_time_s": 18000.0,
        "stable_air": stable,
        "temperature_range_c": [
            float(np.min(simulation.temperature_c)),
            float(np.max(simulation.temperature_c)),
        ],
        "moisture_range_kg_per_kg": [
            float(np.min(simulation.moisture)),
            float(np.max(simulation.moisture)),
        ],
        "maximum_moisture_start": float(maximum_series[0]),
        "maximum_moisture_end": float(maximum_series[-1]),
        "maximum_recorded_increase": maximum_increase,
        "max_coupling_iterations": simulation.max_coupling_iterations,
        "moisture_balance_relative_error": simulation.moisture_balance_relative_error,
        "max_temperature_linear_residual": simulation.max_temperature_linear_residual,
        "max_moisture_linear_residual": simulation.max_moisture_linear_residual,
    }
    output = RESULTS_DIR / "P1_q3_smoke.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


def run_q3() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    run_started = time.perf_counter()
    maximum_time_s = 7.0 * 24.0 * 3600.0
    stable = stable_air_statistics()
    boundary_temperature = stable["temperature_mean_c"]
    boundary_moisture = stable["moisture_mean_kg_per_kg"]
    common_arguments = {
        "model": "q23",
        "end_time_s": maximum_time_s,
        "intervals": 160,
        "record_interval_s": 60.0,
        "stop_at_dryness": True,
        "temporal_scheme": "bdf2",
        "post_observation_temperature_c": boundary_temperature,
        "post_observation_moisture": boundary_moisture,
    }

    base = simulate_variable_model(time_step_s=1.0, **common_arguments)
    time_half = simulate_variable_model(time_step_s=0.5, **common_arguments)
    time_quarter = simulate_variable_model(time_step_s=0.25, **common_arguments)
    if base.dry_time_s is None or time_half.dry_time_s is None or time_quarter.dry_time_s is None:
        raise RuntimeError("时间步验证缺少烘干结束时刻")
    time_base_difference = abs(base.dry_time_s - time_quarter.dry_time_s)
    time_half_difference = abs(time_half.dry_time_s - time_quarter.dry_time_s)
    adjacent_base_half = abs(base.dry_time_s - time_half.dry_time_s)
    adjacent_half_quarter = time_half_difference
    time_gate_passed = (
        time_base_difference <= 1.0
        and time_half_difference <= 0.5
        and adjacent_half_quarter <= adjacent_base_half + 1.0e-12
    )
    authoritative = base if time_gate_passed else time_quarter
    authoritative_time_step = 1.0 if time_gate_passed else 0.25

    space_refined = simulate_variable_model(
        model="q23",
        end_time_s=maximum_time_s,
        time_step_s=authoritative_time_step,
        intervals=320,
        record_interval_s=60.0,
        stop_at_dryness=True,
        temporal_scheme="bdf2",
        post_observation_temperature_c=boundary_temperature,
        post_observation_moisture=boundary_moisture,
    )
    if authoritative.dry_time_s is None or space_refined.dry_time_s is None:
        raise RuntimeError("空间加密验证缺少烘干结束时刻")
    radii_cm = [0.0, 0.5, 1.0, 1.5, 2.0]
    space_table_difference = q3_common_table_difference(
        authoritative, space_refined, radii_cm
    )
    space_time_difference_s = abs(
        authoritative.dry_time_s - space_refined.dry_time_s
    )
    space_time_relative_difference = (
        space_time_difference_s / space_refined.dry_time_s
    )
    if space_time_relative_difference > 0.001:
        space_fine = simulate_variable_model(
            model="q23",
            end_time_s=maximum_time_s,
            time_step_s=authoritative_time_step,
            intervals=640,
            record_interval_s=60.0,
            stop_at_dryness=True,
            temporal_scheme="bdf2",
            post_observation_temperature_c=boundary_temperature,
            post_observation_moisture=boundary_moisture,
        )
        if space_fine.dry_time_s is None:
            raise RuntimeError("640区间空间复验缺少烘干结束时刻")
        fine_relative_difference = abs(
            space_refined.dry_time_s - space_fine.dry_time_s
        ) / space_fine.dry_time_s
        if fine_relative_difference > 0.001:
            raise RuntimeError("640区间复验后烘干时长仍未通过0.1%空间阈值")
        authoritative = space_fine
        authoritative_intervals = 640
    elif space_table_difference > 5.0e-5:
        authoritative = space_refined
        authoritative_intervals = 320
        fine_relative_difference = None
    else:
        authoritative_intervals = 160
        fine_relative_difference = None

    if authoritative.dry_time_s is None:
        raise RuntimeError("问题三权威解缺少烘干时长")
    maximum_series = np.max(authoritative.moisture, axis=1)
    maximum_recorded_increase = float(np.max(np.diff(maximum_series)))
    if maximum_recorded_increase > 1.0e-8:
        raise RuntimeError("问题三权威解的最大含水率出现异常回升")
    if authoritative.moisture_balance_relative_error > 1.0e-8:
        raise RuntimeError("问题三权威解的水分平衡残差超过合同阈值")
    if authoritative.max_temperature_linear_residual > 1.0e-12:
        raise RuntimeError("问题三权威解的温度代数残差超过合同阈值")
    if authoritative.max_moisture_linear_residual > 1.0e-12:
        raise RuntimeError("问题三权威解的含水率代数残差超过合同阈值")

    write_result3(authoritative, RESULTS_DIR / "result3.xlsx")
    table_times, table = q3_periodic_table(authoritative, radii_cm)
    with (RESULTS_DIR / "问题3_表5含水率.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间_h", *[f"{radius:g}_cm" for radius in radii_cm]])
        for time_s, row in zip(table_times, table, strict=True):
            writer.writerow(
                [
                    f"{time_s / 3600.0:.6f}",
                    *[f"{float(value):.8f}" for value in row],
                ]
            )
    write_q3_process_csv(authoritative, RESULTS_DIR / "问题3_关键过程.csv")
    with (RESULTS_DIR / "问题3_迭代次数.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间_s", "Picard迭代次数"])
        writer.writerows(
            zip(
                authoritative.time_s.tolist(),
                authoritative.coupling_iterations.tolist(),
                strict=True,
            )
        )

    with (RESULTS_DIR / "问题3_时间步敏感性.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间步_s", "烘干时长_s", "烘干时长_h", "相对0.25s绝对差_s"])
        for step, simulation in ((1.0, base), (0.5, time_half), (0.25, time_quarter)):
            writer.writerow(
                [
                    step,
                    f"{float(simulation.dry_time_s):.6f}",
                    f"{float(simulation.dry_time_s) / 3600.0:.10f}",
                    f"{abs(float(simulation.dry_time_s) - time_quarter.dry_time_s):.6f}",
                ]
            )
    with (RESULTS_DIR / "问题3_空间网格敏感性.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        writer = csv.writer(stream)
        writer.writerow(["径向区间数", "时间步_s", "烘干时长_s", "烘干时长_h"])
        writer.writerow([160, authoritative_time_step, authoritative.dry_time_s if authoritative_intervals == 160 else (base.dry_time_s if authoritative_time_step == 1.0 else time_quarter.dry_time_s), (base.dry_time_s if authoritative_time_step == 1.0 else time_quarter.dry_time_s) / 3600.0])
        writer.writerow([320, authoritative_time_step, space_refined.dry_time_s, space_refined.dry_time_s / 3600.0])
        if authoritative_intervals == 640:
            writer.writerow([640, authoritative_time_step, authoritative.dry_time_s, authoritative.dry_time_s / 3600.0])

    sensitivity_cases = {
        "温度_减1总体标准差": (
            boundary_temperature - stable["temperature_std_c"],
            boundary_moisture,
            HEAT_TRANSFER_COEFFICIENT,
            MASS_TRANSFER_COEFFICIENT,
        ),
        "温度_加1总体标准差": (
            boundary_temperature + stable["temperature_std_c"],
            boundary_moisture,
            HEAT_TRANSFER_COEFFICIENT,
            MASS_TRANSFER_COEFFICIENT,
        ),
        "空气水分_减1总体标准差": (
            boundary_temperature,
            boundary_moisture - stable["moisture_std_kg_per_kg"],
            HEAT_TRANSFER_COEFFICIENT,
            MASS_TRANSFER_COEFFICIENT,
        ),
        "空气水分_加1总体标准差": (
            boundary_temperature,
            boundary_moisture + stable["moisture_std_kg_per_kg"],
            HEAT_TRANSFER_COEFFICIENT,
            MASS_TRANSFER_COEFFICIENT,
        ),
        "附件末条记录": (
            stable["last_temperature_c"],
            stable["last_moisture_kg_per_kg"],
            HEAT_TRANSFER_COEFFICIENT,
            MASS_TRANSFER_COEFFICIENT,
        ),
        "端面通量上界": (
            boundary_temperature,
            boundary_moisture,
            1.08 * HEAT_TRANSFER_COEFFICIENT,
            1.08 * MASS_TRANSFER_COEFFICIENT,
        ),
    }
    sensitivity = []
    for case_name, (
        case_temperature,
        case_moisture,
        case_heat_coefficient,
        case_mass_coefficient,
    ) in sensitivity_cases.items():
        case = simulate_variable_model(
            model="q23",
            end_time_s=maximum_time_s,
            time_step_s=authoritative_time_step,
            intervals=authoritative_intervals,
            record_interval_s=3600.0,
            stop_at_dryness=True,
            temporal_scheme="bdf2",
            heat_transfer_coefficient=case_heat_coefficient,
            mass_transfer_coefficient=case_mass_coefficient,
            post_observation_temperature_c=case_temperature,
            post_observation_moisture=case_moisture,
        )
        if case.dry_time_s is None:
            raise RuntimeError(f"敏感性情景{case_name}缺少烘干时长")
        sensitivity.append(
            {
                "case": case_name,
                "temperature_c": case_temperature,
                "air_moisture_kg_per_kg": case_moisture,
                "heat_transfer_coefficient_w_per_m2_k": case_heat_coefficient,
                "mass_transfer_coefficient_m_per_s": case_mass_coefficient,
                "dry_time_s": case.dry_time_s,
                "dry_time_h": case.dry_time_s / 3600.0,
                "relative_change": (
                    case.dry_time_s - authoritative.dry_time_s
                )
                / authoritative.dry_time_s,
            }
        )
    with (RESULTS_DIR / "问题3_边界敏感性.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        fieldnames = list(sensitivity[0].keys())
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(sensitivity)

    end_face_case = next(item for item in sensitivity if item["case"] == "端面通量上界")
    if abs(float(end_face_case["relative_change"])) > 0.08:
        raise RuntimeError("端面通量上界情景超过8%，必须回退二维圆柱模型")
    condition_numbers = final_system_condition_numbers(
        authoritative,
        time_step_s=authoritative_time_step,
        heat_transfer_coefficient=HEAT_TRANSFER_COEFFICIENT,
        mass_transfer_coefficient=MASS_TRANSFER_COEFFICIENT,
        post_observation_temperature_c=boundary_temperature,
        post_observation_moisture=boundary_moisture,
    )
    if max(condition_numbers.values()) > 1.0e8:
        raise RuntimeError("问题三末时刻离散系统条件数超过合同阈值")

    summary = {
        "mode": "q3_full",
        "inputs": {
            str(AIR_PATH.relative_to(PROJECT_ROOT)): sha256_file(AIR_PATH),
            str((TEMPLATE_DIR / "result3.xlsx").relative_to(PROJECT_ROOT)): sha256_file(TEMPLATE_DIR / "result3.xlsx"),
        },
        "stable_air": stable,
        "parameters": {
            "temporal_scheme": "backward_euler_start_then_bdf2",
            "authoritative_time_step_s": authoritative_time_step,
            "authoritative_radial_intervals": authoritative_intervals,
            "record_interval_s": 60.0,
            "drying_threshold_kg_per_kg": DRYING_THRESHOLD,
            "heat_transfer_coefficient_w_per_m2_k": HEAT_TRANSFER_COEFFICIENT,
            "mass_transfer_coefficient_m_per_s": MASS_TRANSFER_COEFFICIENT,
            "coupling_temperature_tolerance_c": 1.0e-8,
            "coupling_moisture_tolerance_kg_per_kg": 1.0e-10,
            "coupling_max_iterations": 40,
        },
        "dry_time_s": authoritative.dry_time_s,
        "dry_time_h": authoritative.dry_time_s / 3600.0,
        "table_times_s": table_times,
        "requested_radii_cm": radii_cm,
        "moisture_kg_per_kg": np.round(table, 8).tolist(),
        "sensitivity": sensitivity,
        "verification": {
            "time_step_gate_passed": time_gate_passed,
            "time_step_dry_times_s": {
                "1.0": base.dry_time_s,
                "0.5": time_half.dry_time_s,
                "0.25": time_quarter.dry_time_s,
            },
            "time_base_to_reference_difference_s": time_base_difference,
            "time_half_to_reference_difference_s": time_half_difference,
            "space_refined_dry_time_s": space_refined.dry_time_s,
            "space_time_difference_s": space_time_difference_s,
            "space_time_relative_difference": space_time_relative_difference,
            "space_table_max_absolute_difference_kg_per_kg": space_table_difference,
            "space_fine_relative_difference": fine_relative_difference,
            "maximum_recorded_moisture_increase_kg_per_kg": maximum_recorded_increase,
            "moisture_balance_relative_error": authoritative.moisture_balance_relative_error,
            "max_coupling_iterations": authoritative.max_coupling_iterations,
            "max_temperature_linear_residual": authoritative.max_temperature_linear_residual,
            "max_moisture_linear_residual": authoritative.max_moisture_linear_residual,
            "final_system_condition_numbers": condition_numbers,
        },
        "runtime_seconds": time.perf_counter() - run_started,
        "environment": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "scipy": __import__("scipy").__version__,
            "openpyxl": openpyxl.__version__,
            "platform": platform.platform(),
        },
    }
    (RESULTS_DIR / "问题3_结果摘要.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


def run_q4() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    maximum_time_s = 7.0 * 24.0 * 3600.0
    authoritative = simulate_variable_model(
        model="q4",
        end_time_s=maximum_time_s,
        time_step_s=5.0,
        intervals=160,
        record_interval_s=60.0,
        stop_at_dryness=True,
    )
    time_coarse = simulate_variable_model(
        model="q4",
        end_time_s=maximum_time_s,
        time_step_s=10.0,
        intervals=160,
        record_interval_s=60.0,
        stop_at_dryness=True,
    )
    space_coarse = simulate_variable_model(
        model="q4",
        end_time_s=maximum_time_s,
        time_step_s=5.0,
        intervals=80,
        record_interval_s=60.0,
        stop_at_dryness=True,
    )
    write_result4(authoritative, RESULTS_DIR / "result4.xlsx")
    if authoritative.dry_time_s is None:
        raise RuntimeError("问题四未返回干燥时长")
    six_hour_times = list(
        np.arange(6.0 * 3600.0, authoritative.dry_time_s, 6.0 * 3600.0)
    )
    table_times = six_hour_times + [authoritative.dry_time_s]
    summary_rows = []
    for time_s in table_times:
        index = int(np.argmin(np.abs(authoritative.time_s - time_s)))
        current_radius_cm = 100.0 * authoritative.radius_m[index]
        requested = [
            radius for radius in (0.0, 0.5, 1.0, 1.5) if radius < current_radius_cm
        ]
        values = interpolate_at_physical_radii(
            authoritative, authoritative.moisture, index, requested
        ).tolist()
        summary_rows.append(
            {
                "time_s": time_s,
                "radius_cm": current_radius_cm,
                "fixed_radii_cm": requested,
                "fixed_values": values,
                "surface_value": float(authoritative.moisture[index, -1]),
            }
        )
    summary = {
        "mode": "q4_full",
        "inputs": {
            str(AIR_PATH.relative_to(PROJECT_ROOT)): sha256_file(AIR_PATH),
            str(RADIUS_PATH.relative_to(PROJECT_ROOT)): sha256_file(RADIUS_PATH),
        },
        "dry_time_s": authoritative.dry_time_s,
        "dry_time_h": authoritative.dry_time_s / 3600.0,
        "table": summary_rows,
        "verification": {
            "time_coarse_dry_time_s": time_coarse.dry_time_s,
            "space_coarse_dry_time_s": space_coarse.dry_time_s,
            "time_step_difference_s": float(
                abs(authoritative.dry_time_s - float(time_coarse.dry_time_s))
            ),
            "space_grid_difference_s": float(
                abs(authoritative.dry_time_s - float(space_coarse.dry_time_s))
            ),
            "moisture_balance_relative_error": authoritative.moisture_balance_relative_error,
            "max_coupling_iterations": authoritative.max_coupling_iterations,
        },
    }
    (RESULTS_DIR / "问题4_结果摘要.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser(description="2026 CUMCM A题全程干燥求解")
    parser.add_argument(
        "--mode",
        choices=("smoke", "q3-smoke", "q2", "q3", "q4", "all"),
        default="smoke",
    )
    arguments = parser.parse_args()
    if arguments.mode == "smoke":
        run_smoke()
    elif arguments.mode == "q3-smoke":
        run_q3_smoke()
    elif arguments.mode == "q2":
        run_q2()
    elif arguments.mode == "q3":
        run_q3()
    elif arguments.mode == "q4":
        run_q4()
    else:
        run_q2()
        run_q3()
        run_q4()


if __name__ == "__main__":
    main()
