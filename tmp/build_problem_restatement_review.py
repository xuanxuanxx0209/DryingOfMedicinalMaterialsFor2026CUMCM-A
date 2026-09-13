from pathlib import Path
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


PROJECT_ROOT = Path(r"C:\Users\27632\Desktop\26CUMCM_A")
DOCX_SCRIPTS = Path(
    r"C:\Users\27632\.codex\skills\cumcm-step-review\tools\docx\scripts"
)
sys.path.insert(0, str(DOCX_SCRIPTS))

import paper_format as pf  # noqa: E402


def set_run(run, *, font="宋体", size=12, bold=False):
    pf.set_run_font(run, font=font, size=size, bold=bold)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return run


def add_body(doc, text, *, lead=None):
    paragraph = pf.paragraph(doc, first_line=True, line_spacing=1.25)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.widow_control = True
    if lead:
        set_run(paragraph.add_run(lead), bold=True)
    set_run(paragraph.add_run(text))
    return paragraph


def add_reference(doc, text):
    paragraph = pf.paragraph(doc, line_spacing=1.15)
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(-24)
    paragraph.paragraph_format.space_after = Pt(3)
    set_run(paragraph.add_run(text), size=10.5)
    return paragraph


def add_page_number(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._element.append(element)
    set_run(run, size=10.5)


def remove_style_borders(style):
    properties = style.element.find(qn("w:pPr"))
    if properties is None:
        return
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


doc = pf.new_document(contest="cumcm")
doc.core_properties.title = "收缩圆柱药材烘干的变物性热质传递建模"
doc.core_properties.subject = "全国大学生数学建模竞赛论文草稿"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
remove_style_borders(doc.styles["Title"])

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(12)
set_run(title.add_run("收缩圆柱药材烘干的变物性热质传递建模"), font="黑体", size=16)

section_title = pf.paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
section_title.paragraph_format.space_before = Pt(6)
section_title.paragraph_format.space_after = Pt(8)
set_run(section_title.add_run("1 问题重述"), font="黑体", size=16, bold=True)
pf.heading2(doc, "1.1 问题背景")
add_body(
    doc,
    "中药材热风干燥由表面对流换热、传质与药材内部的热量传导、水分扩散共同驱动，通常经历预热平衡和恒温干燥两个阶段。烘干过程中，烘房温度与空气水分浓度随时间变化，药材的热物性和水分扩散能力又受温度、含水率影响，持续失水还会引起外形尺寸收缩。因此，仅凭表面状态难以判断药材内部各处的温度和干基含水率是否均达到工艺要求。圆柱物料的热风干燥具有典型的非稳态热质传递特征[1,2]。题目据此给出烘房环境记录、药材物性关系及半径变化数据，要求刻画内部温度场与含水率场的时空演化，并依据全域含水率约束确定烘干终点。",
)

pf.heading2(doc, "1.2 问题要求")
add_body(doc, "本文需要解决以下四个问题：")
add_body(
    doc,
    "以长 25 cm、初始半径 2 cm 的圆柱形药材为研究对象，根据给定的烘房温度、空气水分浓度和热质传递参数，建立预热平衡阶段的温度与干基含水率变化模型。药材初始温度为 28 ℃，初始干基含水率为 2.55 kg/kg。除计算指定时刻和指定径向位置的结果外，还需按 1 s 的时间间隔、0.1 cm 的空间间隔输出前 1800 s 的完整结果。",
    lead="问题一：",
)
add_body(
    doc,
    "在问题一的基础上，区分预热平衡与恒温干燥阶段的参数，并考虑药材物性随温度和含水率的变化，建立约 2～3 天全程烘干过程中温度与干基含水率的变化模型。给出前 3 h 内每隔 0.5 h、距轴线每隔 0.5 cm 的计算结果，并按 1 s 的时间间隔、0.1 cm 的空间间隔输出完整结果。",
    lead="问题二：",
)
add_body(
    doc,
    "沿用问题二的全程烘干模型，在药材尺寸保持不变的条件下，以各处干基含水率均严格低于 0.15 kg/kg 作为结束判据，确定烘干时长。给出每隔 6 h、距轴线每隔 0.5 cm 的含水率，并按 60 s 的时间间隔、0.1 cm 的空间间隔输出完整结果。",
    lead="问题三：",
)
add_body(
    doc,
    "进一步考虑失水引起的尺寸变化，将实测半径变化及相应的物性关系纳入模型，沿用问题三的结束判据重新确定烘干时长。给出每隔 6 h、距轴线每隔 0.5 cm 以及药材表面的含水率，并按 60 s 的时间间隔、0.1 cm 的空间间隔输出完整结果。",
    lead="问题四：",
)

pf.heading2(doc, "参考文献")
add_reference(
    doc,
    "[1] 刘格含, 王鹏, 吴小华, 等. 农产品热风干燥传热传质数值模拟研究进展[J]. 食品工业科技, 2020, 41(22): 342-350, 357.",
)
add_reference(
    doc,
    "[2] TZEMPELIKOS D A, MITRAKOS D, VOUROS A P, et al. Numerical modeling of heat and mass transfer during convective drying of cylindrical quince slices[J]. Journal of Food Engineering, 2015, 156: 10-21.",
)

for section in doc.sections:
    add_page_number(section)

output = PROJECT_ROOT / "论文草稿.docx"
temporary = output.with_name(f".{output.name}.tmp")
doc.save(temporary)
temporary.replace(output)
print(output)
