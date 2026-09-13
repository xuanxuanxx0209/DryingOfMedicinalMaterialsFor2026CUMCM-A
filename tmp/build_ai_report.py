from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "AI使用报告.docx"


def set_font(run, name="宋体", size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_cell_text(cell, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_font(r, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")


def shade(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for j, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row_data):
            set_cell_text(cells[j], str(value), align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    if widths:
        for row in table.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        set_font(run, name="黑体", size=14 if level == 1 else 12, bold=True)
    return p


def body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(5)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2)
    else:
        set_font(p.add_run(text))
    return p


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("全国大学生数学建模竞赛"), name="黑体", size=18, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(18)
    set_font(p2.add_run("AI 工具使用详情"), name="黑体", size=22, bold=True)

    heading(doc, "一 论文中的 AI 工具使用声明")
    add_table(
        doc,
        ["选择", "声明内容"],
        [["√", "本参赛队在竞赛过程中使用了 AI 工具，主要用于程序调试、图表制作、论文结构起草、语言校核及排版，详细使用情况见支撑材料。"],
         ["□", "本参赛队在竞赛过程中未使用任何 AI 工具。"]],
        [2.2, 13.6],
    )

    heading(doc, "二 所用 AI 工具清单")
    add_table(
        doc,
        ["序号", "工具名称", "版本或型号", "提供方或入口", "用途"],
        [["1", "Codex", "当前会话模型，版本由平台管理", "OpenAI Codex 桌面应用", "程序调试、图表制作、论文结构起草、语言校核及排版"]],
        [1.2, 2.0, 4.3, 3.2, 5.1],
    )

    heading(doc, "第一部分 使用概述")
    body(doc, "本次 AI 工具使用集中在程序调试、图表制作、论文结构起草、语言校核及排版五类辅助工作。使用内容不包含以 AI 输出替代程序实际运行结果，论文中的数据、表格和图件均以项目文件及计算输出为准。")
    add_table(
        doc,
        ["使用方面", "主要目的", "对应文件或位置"],
        [
            ["程序调试", "检查路径、依赖、语法、输入输出和运行报错，协助定位并修正程序问题", "各问题求解与绘图程序；支撑材料"],
            ["图表制作", "依据已有结果数据生成论文图件，统一字体、配色、尺寸和导出格式", "问题1_绘图.py 至问题4_绘图.py；figures"],
            ["论文结构起草", "整理章节顺序、标题层级、参考文献和附录文件清单", "完整论文-LaTeX/example.tex"],
            ["语言校核", "检查术语、符号、单位、图表引用和语句衔接的一致性", "论文正文、参考文献及图表说明"],
            ["排版", "协助处理 LaTeX 表格、代码附录、分页和 PDF 编译检查", "完整论文-LaTeX；完整论文.pdf"],
        ],
        [3.0, 7.0, 6.2],
    )

    heading(doc, "第二部分 分环节使用记录")
    records = [
        ("2.1 程序调试", "2026-09-13", "检查各脚本的本机硬编码路径、模块依赖和命令行参数；根据报错定位问题并提出修改；运行语法编译、导入检查和四项冒烟测试。", "检查项目代码的路径与依赖，修正影响异机运行的问题，并验证四问程序能够执行。", "相关修改写入求解、复现和绘图脚本；最终以实际运行输出确认程序状态。", "各问题求解程序、问题4_复现.py、支撑材料/results 中的测试输出。"),
        ("2.2 图表制作", "2026-09-13", "依据已有 CSV、JSON 和工作簿结果协助生成论文图件，统一字体、配色、线宽、坐标单位、图例和 PDF 导出规格，并检查图件与正文引用的对应关系。", "根据项目结果生成适合国赛论文的图件，保持数据不变并统一图形样式。", "图件只读取项目结果文件；样式和版面问题修改后重新导出。", "问题1_绘图.py 至问题4_绘图.py、论文图件优化.py、figures 目录。"),
        ("2.3 论文结构起草", "2026-09-13", "协助整理章节层级和内容顺序，补充 AI 工具使用声明、参考文献、支撑材料清单及代码附录，并生成仅正文与完整附录两种编译入口。", "根据现有论文和项目文件完善论文末尾结构，同时保留正文版和含附录版。", "结构建议写入 LaTeX 源文件；未改变模型计算结果和正文关键数值。", "完整论文-LaTeX/example.tex、example-body.tex。"),
        ("2.4 语言校核", "2026-09-13", "检查语句通顺性、术语与符号一致性、单位写法、图表引用和参考文献格式，对重复、歧义或不够规范的表达提出修改。", "校核论文语言和术语，保持原有模型含义、数据与结论不变。", "修改限于表达和一致性问题，不据此生成新的计算结论。", "论文正文、图表说明、AI 工具使用声明与参考文献。"),
        ("2.5 排版", "2026-09-13", "协助调整 LaTeX 表格、长表格、代码清单和分页逻辑，检查 AI 声明位置、参考文献顺序、附录起始页、PDF 页面尺寸及压缩包文件清单。", "完成论文排版和编译检查，分别输出仅正文版和含附录版 PDF。", "通过实际编译和页面检查处理排版问题。", "完整论文-仅正文.pdf、完整论文.pdf、支撑材料.zip。"),
    ]
    for title, when, output, prompt, action, verify in records:
        if title.startswith("2.4 "):
            doc.add_page_break()
        heading(doc, title, 2)
        add_table(doc, ["条目", "填写内容"], [
            ["使用时间", when],
            ["工具及版本", "Codex；当前会话模型，版本由平台管理"],
            ["提示词摘要", prompt],
            ["使用过程", output],
            ["输出使用情况", action],
            ["对应位置", verify],
        ], [3.4, 12.4])

    heading(doc, "第三部分 提示方式与使用过程")
    body(doc, "提示内容以任务目标、项目文件和运行现象为依据，以下为脱敏后的概括，不包含账号、身份信息及无关对话。")
    add_table(doc, ["方式", "是否使用", "说明"], [
        ["任务说明", "是", "说明需要调试的程序、图表目标、论文结构或排版要求"],
        ["文件检查", "是", "读取相关代码、结果文件和 LaTeX 源文件以定位问题"],
        ["报错与结果反馈", "是", "根据实际运行报错、编译日志或图件效果继续调整"],
        ["迭代修改", "是", "对修改后的程序、图表和文档再次运行或编译检查"],
    ], [4.0, 2.2, 9.6])
    heading(doc, "3.1 典型使用示例", 2)
    body(doc, "程序调试示例：提示 AI 检查项目代码中的硬编码路径和外部依赖，并根据实际运行结果修正；图表制作示例：提示 AI 依据既有结果文件统一图件字体、配色、尺寸和导出格式；论文处理示例：提示 AI 补充论文末尾结构、校核语言与引用，并分别编译仅正文版和含附录版 PDF。")

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("AI 工具使用详情"), size=9)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "全国大学生数学建模竞赛 AI 工具使用详情"
    zoom = doc.settings._element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
